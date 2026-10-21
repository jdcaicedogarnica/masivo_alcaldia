import os
import shutil
import re
import datetime
import openpyxl
import warnings
import time
from openpyxl.styles import Font
from pypdf import PdfReader
from tkinter import filedialog, Tk

# Silenciamos advertencias de estilos de Excel
warnings.filterwarnings("ignore", category=UserWarning, module='openpyxl')

def seleccionar_carpeta(titulo):
    """Abre ventana para seleccionar carpetas asegurando que salte al frente"""
    print(f"\n[ESPERANDO] {titulo}...")
    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)  # Fuerza a la ventana a estar encima de todo
    carpeta = filedialog.askdirectory(title=titulo)
    root.update()
    root.destroy()
    if carpeta:
        print(f" -> Seleccionado: {carpeta}")
    return carpeta

def obtener_nombre_limpio(nombre_archivo):
    """Quita la extension .pdf y espacios extra"""
    return os.path.splitext(nombre_archivo)[0].strip()

def analizar_pdf(ruta_pdf):
    """
    Lee la primera pagina del PDF para encontrar el encabezado SHM.
    Calcula peso en MB y cuenta paginas.
    """
    info = {
        "paginas": 0,
        "radicado": "NO ENCONTRADO",
        "tamano_mb": "0 MB",
        "fecha": "01/01/2000"
    }

    # 1. CALCULO DE PESO (MB) Y FECHA
    try:
        stats = os.stat(ruta_pdf)
        mb_size = round(stats.st_size / (1024 * 1024), 2)
        info['tamano_mb'] = f"{mb_size} MB"
        fecha_dt = datetime.datetime.fromtimestamp(stats.st_mtime)
        info['fecha'] = fecha_dt.strftime("%d/%m/%Y")
    except:
        pass

    # 2. LECTURA DE TEXTO INTERNO (BUSQUEDA DE SHM)
    try:
        reader = PdfReader(ruta_pdf)
        info['paginas'] = len(reader.pages)
        if len(reader.pages) > 0:
            texto = reader.pages[0].extract_text()
            texto_lineal = texto.replace('\n', ' ')
            # Busca el patrón SHM específico
            patron = r'(SHM\s*[-]?\s*\d+\s*[-]?\s*\d+)'
            match = re.search(patron, texto_lineal, re.IGNORECASE)
            if match:
                info['radicado'] = match.group(1).replace(' ', '')
    except:
        pass

    return info

def iniciar_proceso_final():
    print("====================================================")
    print("   SISTEMA DE ORGANIZACION DE COBRO COACTIVO")
    print("====================================================")
    print("Por favor, atienda las ventanas de seleccion de carpeta.")

    # --- SELECCION DE RUTAS ---
    ruta_origen = seleccionar_carpeta("1. CARPETA DE ORIGEN (DONDE ESTAN LOS PDF)")
    if not ruta_origen: return

    ruta_plantilla_dir = seleccionar_carpeta("2. CARPETA QUE CONTIENE LA PLANTILLA (.xlsm)")
    if not ruta_plantilla_dir: return
    
    nombre_plantilla = '00FormatoIndiceElectronico.xlsm'
    ruta_plantilla_full = os.path.join(ruta_plantilla_dir, nombre_plantilla)
    
    if not os.path.exists(ruta_plantilla_full):
        print(f"\n[ERROR] No se encontro '{nombre_plantilla}' en la carpeta elegida.")
        return

    ruta_destino = seleccionar_carpeta("3. CARPETA DE DESTINO FINAL PARA LOS EXPEDIENTES")
    if not ruta_destino: return

    # --- PROCESAMIENTO ---
    try:
        archivos = [f for f in os.listdir(ruta_origen) if f.lower().endswith('.pdf')]
    except Exception as e:
        print(f"Error al acceder a la carpeta: {e}")
        return

    print(f"\n[PROCESANDO] {len(archivos)} archivos encontrados.\n")

    for archivo in archivos:
        path_pdf = os.path.join(ruta_origen, archivo)
        nombre_contribuyente = obtener_nombre_limpio(archivo)
        
        datos = analizar_pdf(path_pdf)
        
        print(f" [+] Contribuyente: {nombre_contribuyente}")
        
        # Crear Carpeta Individual
        carpeta_final = os.path.join(ruta_destino, nombre_contribuyente)
        os.makedirs(carpeta_final, exist_ok=True)

        # Copiar PDF original a la carpeta del contribuyente
        shutil.copy2(path_pdf, os.path.join(carpeta_final, archivo))

        # Copiar y Llenar Plantilla Excel
        dest_excel = os.path.join(carpeta_final, nombre_plantilla)
        try:
            shutil.copy2(ruta_plantilla_full, dest_excel)
            wb = openpyxl.load_workbook(dest_excel, keep_vba=True)
            ws = wb.active
            
            # --- MAPEO DE CELDAS ACTUALIZADO ---
            ws['B5'] = nombre_contribuyente    # Nombre del contribuyente
            ws['B10'] = datos['fecha']         # Fecha inicio
            ws['C10'] = datos['fecha']         # Fecha fin
            ws['E10'] = datos['paginas']       # Numero de paginas
            
            ws['G10'] = datos['tamano_mb']     # TAMAÑO DEL ARCHIVO EN G10
            ws['I10'] = datos['radicado']      # RADICADO SHM EN I10 (Cambio solicitado)
            
            wb.save(dest_excel)
            wb.close()
            print(f"     -> Exito: Radicado {datos['radicado']} | Peso {datos['tamano_mb']}")
            
        except Exception as e:
            print(f"     [!] Error procesando Excel de {nombre_contribuyente}: {e}")

    print("\n====================================================")
    print("      PROCESO FINALIZADO EXITOSAMENTE")
    print("====================================================")
    time.sleep(3)

if __name__ == "__main__":
    iniciar_proceso_final()